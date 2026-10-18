"""Lector universal de archivos (0.1.2): lo que adjuntas en el chat o mandas por WhatsApp.

read_any(datos, nombre) -> {"kind", "label", "text", "truncated", "note", ...} o {"error"}.

  texto y código (cualquier codificación) · PDF · Word (.docx) · Excel (.xlsx, .xlsm) · CSV ·
  PowerPoint (.pptx) · OpenDocument (.odt, .ods, .odp) · EPUB · RTF · HTML/XML · correo (.eml) ·
  imágenes (visión local o Gemini) · audio y vídeo (transcripción con Whisper, hasta 20 min) ·
  ZIP (lista y lee lo que hay dentro) · programas (.exe, .dll, .msi: se identifican, NUNCA se ejecutan) ·
  cualquier otro binario (se dice qué es y su huella).

Nada se ejecuta ni se abre con otros programas: todo se lee en memoria. El texto que sale de un archivo
es dato no confiable (quien llama lo marca con safety.mark_untrusted).
"""

import email
import hashlib
import io
import mimetypes
import re
import zipfile
from email import policy
from html import unescape
from pathlib import Path
from typing import Any, Dict, List, Optional

MAX_TEXT = 60_000          # texto extraído que se guarda (el modelo recibe un recorte)
MAX_MEDIA_SECONDS = 20 * 60
MAX_ZIP_ENTRIES = 200
MAX_ZIP_READ = 5           # archivos de dentro del ZIP que se leen
MAX_ZIP_INNER_BYTES = 10 * 1024 * 1024

TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".log", ".ini", ".cfg", ".conf", ".toml",
            ".yaml", ".yml", ".xml", ".html", ".htm", ".css", ".js", ".mjs", ".ts", ".tsx", ".jsx", ".py", ".ipynb",
            ".java", ".kt", ".c", ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".rb", ".php", ".swift", ".sql", ".sh",
            ".bat", ".cmd", ".ps1", ".r", ".m", ".tex", ".bib", ".srt", ".vtt", ".svg", ".env.example", ".gitignore"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".heic", ".ico"}
AUDIO_EXT = {".mp3", ".wav", ".ogg", ".opus", ".m4a", ".aac", ".flac", ".wma", ".amr", ".weba"}
VIDEO_EXT = {".mp4", ".mkv", ".mov", ".avi", ".webm", ".wmv", ".m4v", ".3gp", ".flv"}
PROGRAM_EXT = {".exe", ".dll", ".msi", ".scr", ".com", ".sys", ".lnk", ".jar", ".apk", ".app", ".dmg"}


def safe_filename(name: Optional[str], default: str = "archivo") -> str:
    clean = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", (name or "").strip()).strip(" .")
    return clean[:120] or default


def _strip_tags(markup: str) -> str:
    markup = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", markup, flags=re.DOTALL | re.IGNORECASE)
    markup = re.sub(r"<(br|/p|/div|/li|/h[1-6]|/tr|text:p|/text:p|/text:h|w:p|/w:p|a:p|/a:p)\b[^>]*>", "\n", markup,
                    flags=re.IGNORECASE)
    text = unescape(re.sub(r"<[^>]+>", "", markup))
    return re.sub(r"\n\s*\n+", "\n\n", re.sub(r"[ \t]+", " ", text)).strip()


def _decode_text(data: bytes) -> Optional[str]:
    """Texto en UTF-8 / UTF-16 / Windows-1252; None si parece binario."""
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace")
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    head = data[:8192]
    if b"\x00" in head:
        return None
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _sniff(data: bytes, ext: str, mime: str) -> str:
    """Tipo por el contenido (manda sobre la extensión, que puede mentir)."""
    h = data[:16]
    if h.startswith(b"%PDF"):
        return "pdf"
    if h.startswith(b"MZ"):
        return "program"
    if h.startswith(b"PK\x03\x04") or h.startswith(b"PK\x05\x06"):
        return "zipfamily"
    if h.startswith(b"\x89PNG") or h.startswith(b"\xff\xd8\xff") or h[:4] in (b"GIF8",) or h.startswith(b"BM") \
            or (h[:4] == b"RIFF" and data[8:12] == b"WEBP"):
        return "image"
    if h.startswith(b"{\\rtf"):
        return "rtf"
    if h[:4] == b"RIFF" and data[8:12] == b"WAVE" or h.startswith(b"ID3") or h.startswith(b"OggS") \
            or h.startswith(b"fLaC") or h[:2] in (b"\xff\xfb", b"\xff\xf3", b"\xff\xf2"):
        return "audio"
    if data[4:8] == b"ftyp":
        return "audio" if ext in AUDIO_EXT else "video"
    if h.startswith(b"\x1aE\xdf\xa3"):  # Matroska / WebM
        return "audio" if ext in (".weba",) else "video"
    if ext in IMAGE_EXT or mime.startswith("image/"):
        return "image"
    if ext in AUDIO_EXT or mime.startswith("audio/"):
        return "audio"
    if ext in VIDEO_EXT or mime.startswith("video/"):
        return "video"
    if ext in PROGRAM_EXT:
        return "program"
    if ext == ".eml" or mime == "message/rfc822":
        return "email"
    return "other"


# ---------------------------------------------------------------- lectores
def _pdf(data: bytes) -> Dict[str, Any]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            return {"error": "El PDF tiene contraseña: no puedo abrirlo."}
    pages = []
    for page in reader.pages[:200]:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    text = "\n\n".join(f"[Página {i}]\n{p.strip()}" for i, p in enumerate(pages, 1) if p.strip())
    if not text:
        return {"error": "El PDF no tiene texto seleccionable (parece escaneado): mándamelo como imagen y lo leo con visión."}
    return {"kind": "pdf", "label": f"PDF de {len(reader.pages)} páginas", "text": text}


def _xml_from_zip(z: zipfile.ZipFile, name: str) -> str:
    return z.read(name).decode("utf-8", errors="replace")


def _docx(z: zipfile.ZipFile) -> Dict[str, Any]:
    xml = _xml_from_zip(z, "word/document.xml")
    xml = re.sub(r"<w:tab/>", "\t", xml)
    return {"kind": "word", "label": "documento de Word", "text": _strip_tags(xml)}


def _pptx(z: zipfile.ZipFile) -> Dict[str, Any]:
    slides = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                    key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)))
    parts = []
    for i, name in enumerate(slides, 1):
        texts = re.findall(r"<a:t>([^<]*)</a:t>", _xml_from_zip(z, name))
        if texts:
            parts.append(f"[Diapositiva {i}]\n" + "\n".join(unescape(t) for t in texts))
    return {"kind": "presentacion", "label": f"presentación de {len(slides)} diapositivas", "text": "\n\n".join(parts)}


def _xlsx(data: bytes) -> Dict[str, Any]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets[:20]:
        rows = []
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            if r >= 500:
                rows.append("… (más filas)")
                break
            cells = ["" if v is None else str(v) for v in row]
            while cells and cells[-1] == "":
                cells.pop()
            if cells:
                rows.append("\t".join(cells))
        parts.append(f"[Hoja «{ws.title}»]\n" + ("\n".join(rows) if rows else "(vacía)"))
    wb.close()
    return {"kind": "hoja", "label": f"hoja de cálculo ({len(wb.worksheets)} hojas)", "text": "\n\n".join(parts)}


def _opendocument(z: zipfile.ZipFile, mimetype: str) -> Dict[str, Any]:
    kind = {"text": ("word", "documento de OpenDocument"), "spreadsheet": ("hoja", "hoja de OpenDocument"),
            "presentation": ("presentacion", "presentación de OpenDocument")}
    key = next((k for k in kind if k in mimetype), "text")
    xml = _xml_from_zip(z, "content.xml")
    xml = re.sub(r"</table:table-cell>", "\t", xml)
    xml = re.sub(r"</table:table-row>", "\n", xml)
    return {"kind": kind[key][0], "label": kind[key][1], "text": _strip_tags(xml)}


def _epub(z: zipfile.ZipFile) -> Dict[str, Any]:
    pages = sorted(n for n in z.namelist() if n.lower().endswith((".xhtml", ".html", ".htm")))
    text = "\n\n".join(_strip_tags(_xml_from_zip(z, n)) for n in pages[:300])
    return {"kind": "libro", "label": f"libro EPUB ({len(pages)} secciones)", "text": text}


def _zip(z: zipfile.ZipFile, depth: int) -> Dict[str, Any]:
    infos = [i for i in z.infolist() if not i.is_dir()]
    listing = [f"- {i.filename} ({i.file_size // 1024} KB)" for i in infos[:MAX_ZIP_ENTRIES]]
    more = f"\n… y {len(infos) - MAX_ZIP_ENTRIES} más" if len(infos) > MAX_ZIP_ENTRIES else ""
    parts = [f"Contiene {len(infos)} archivos:\n" + "\n".join(listing) + more]
    read = 0
    if depth < 1:
        for info in infos:
            if read >= MAX_ZIP_READ:
                break
            if info.file_size > MAX_ZIP_INNER_BYTES or info.flag_bits & 0x1:  # grande o con contraseña
                continue
            inner = read_any(z.read(info), Path(info.filename).name, _depth=depth + 1, transcribe=False, see=False)
            if inner.get("text"):
                parts.append(f"[Dentro: {info.filename} · {inner.get('label', '')}]\n{inner['text'][:4000]}")
                read += 1
    return {"kind": "comprimido", "label": f"ZIP con {len(infos)} archivos", "text": "\n\n".join(parts)}


def _zipfamily(data: bytes, ext: str, depth: int) -> Dict[str, Any]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = set(z.namelist())
        if "word/document.xml" in names:
            return _docx(z)
        if "xl/workbook.xml" in names:
            return _xlsx(data)
        if any(n.startswith("ppt/slides/") for n in names):
            return _pptx(z)
        if "mimetype" in names:
            mt = z.read("mimetype").decode("ascii", errors="replace")
            if "opendocument" in mt and "content.xml" in names:
                return _opendocument(z, mt)
            if "epub" in mt:
                return _epub(z)
        return _zip(z, depth)


def _rtf(data: bytes) -> Dict[str, Any]:
    text = data.decode("cp1252", errors="replace")
    text = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: bytes([int(m.group(1), 16)]).decode("cp1252", errors="replace"), text)
    text = re.sub(r"\\u(-?\d+)\??", lambda m: chr(int(m.group(1)) % 65536), text)
    text = re.sub(r"\\(par|line)\b ?", "\n", text)
    text = re.sub(r"\{\\\*[^{}]*\}|\\[a-zA-Z]+-?\d* ?|[{}]", "", text)
    return {"kind": "word", "label": "documento RTF", "text": re.sub(r"\n\s*\n+", "\n\n", text).strip()}


def _email(data: bytes) -> Dict[str, Any]:
    from .tools.email_tools import _text_of  # mismo conversor HTML -> texto que el correo IMAP
    msg = email.message_from_bytes(data, policy=policy.default)
    attachments = [p.get_filename() for p in msg.walk() if p.get_filename()]
    head = f"De: {msg.get('From', '')}\nPara: {msg.get('To', '')}\nAsunto: {msg.get('Subject', '')}\nFecha: {msg.get('Date', '')}"
    if attachments:
        head += "\nAdjuntos: " + ", ".join(attachments)
    return {"kind": "correo", "label": "correo (.eml)", "text": head + "\n\n" + _text_of(msg)}


def _media(data: bytes, kind: str, transcribe: bool) -> Dict[str, Any]:
    import av
    import numpy as np
    label = "audio" if kind == "audio" else "vídeo"
    try:
        with av.open(io.BytesIO(data)) as src:
            duration = float(src.duration / av.time_base) if src.duration else None
            has_audio = any(s.type == "audio" for s in src.streams)
            video = next((s for s in src.streams if s.type == "video"), None)
            size = f" {video.codec_context.width}×{video.codec_context.height}" if video and video.codec_context.width else ""
            if not has_audio:
                return {"kind": kind, "label": f"{label}{size} sin sonido", "text": "",
                        "note": "No tiene pista de audio: no hay nada que transcribir."}
            if not transcribe:
                return {"kind": kind, "label": f"{label}{size}", "text": "", "note": "No se transcribió."}
            chunks, total = [], 0
            resampler = av.AudioResampler(format="flt", layout="mono", rate=16000)
            for frame in src.decode(audio=0):
                for rf in resampler.resample(frame):
                    arr = rf.to_ndarray().reshape(-1)
                    chunks.append(arr)
                    total += arr.size
                if total >= MAX_MEDIA_SECONDS * 16000:
                    break
    except Exception as e:
        return {"error": f"No pude abrir ese {label} ({type(e).__name__})."}
    samples = np.concatenate(chunks).astype(np.float32) if chunks else np.zeros(0, dtype=np.float32)
    mins = f"{duration / 60:.1f} min" if duration else "duración desconocida"
    if samples.size < 1600:
        return {"kind": kind, "label": f"{label} ({mins})", "text": "", "note": "Es demasiado corto o no tiene voz."}
    from .channels.voice_notes import _whisper
    segments, _ = _whisper().transcribe(samples, vad_filter=True, beam_size=1)
    text = " ".join(s.text.strip() for s in segments).strip()
    cut = duration and duration > MAX_MEDIA_SECONDS
    return {"kind": kind, "label": f"{label}{size} ({mins})", "text": text,
            "note": ("Transcripción de los primeros 20 minutos." if cut else "Transcripción automática (Whisper).")
            + ("" if text else " No se entendió ninguna voz.")}


def _image(data: bytes, mime: str, question: str, see: bool) -> Dict[str, Any]:
    if not see:
        return {"kind": "imagen", "label": "imagen", "text": "", "note": "Imagen (no se analizó)."}
    from . import vision
    desc, who = vision.describe_image(data, question or "Describe la imagen con detalle y transcribe todo el texto que tenga.",
                                      mime or "image/png")
    if not desc:
        return {"kind": "imagen", "label": "imagen", "text": "", "note": f"No puedo verla todavía: {who}."}
    return {"kind": "imagen", "label": "imagen", "text": desc, "note": f"Descrita por {who.split(':')[0]}."}


def read_any(data: bytes, filename: str = "", mime: str = "", question: str = "", *, transcribe: bool = True,
             see: bool = True, _depth: int = 0) -> Dict[str, Any]:
    name = safe_filename(filename)
    ext = Path(name).suffix.lower()
    mime = (mime or mimetypes.guess_type(name)[0] or "").lower()
    kind = _sniff(data, ext, mime)
    digest = hashlib.sha256(data).hexdigest()[:16]
    try:
        if kind == "pdf":
            out = _pdf(data)
        elif kind == "zipfamily":
            out = _zipfamily(data, ext, _depth)
        elif kind == "image":
            out = _image(data, mime, question, see)
        elif kind in ("audio", "video"):
            out = _media(data, kind, transcribe)
        elif kind == "rtf":
            out = _rtf(data)
        elif kind == "email":
            out = _email(data)
        elif kind == "program":
            out = {"kind": "programa", "label": "programa o archivo ejecutable", "text": "",
                   "note": f"Es un programa ({ext or 'binario MZ'}). Por seguridad no lo ejecuto ni lo abro. Huella SHA-256: {digest}…"}
        else:
            text = _decode_text(data)
            if text is None:
                if ext == ".xls":
                    return {"kind": "hoja", "label": "hoja de Excel antigua (.xls)", "text": "",
                            "note": "Formato .xls antiguo: guárdalo como .xlsx (Archivo › Guardar como) y lo leo."}
                out = {"kind": "binario", "label": f"archivo binario {ext or ''}".strip(), "text": "",
                       "note": f"No es texto ni un formato que sepa leer ({mime or 'tipo desconocido'}). "
                               f"Huella SHA-256: {digest}…"}
            else:
                if ext in (".html", ".htm") or mime == "text/html":
                    text = _strip_tags(text)
                label = {".csv": "tabla CSV", ".tsv": "tabla TSV", ".json": "JSON", ".md": "Markdown"}.get(ext)
                if not label:
                    label = f"código {ext[1:]}" if ext in TEXT_EXT and ext not in (".txt", ".log") else "texto"
                out = {"kind": "texto", "label": label, "text": text}
    except zipfile.BadZipFile:
        return {"error": "El archivo comprimido está dañado."}
    except Exception as e:
        return {"error": f"No pude leer «{name}» ({type(e).__name__}: {str(e)[:120]})."}
    if "error" in out:
        return out
    text = out.get("text") or ""
    out.update(name=name, size=len(data), sha256=digest, truncated=len(text) > MAX_TEXT, text=text[:MAX_TEXT])
    out.setdefault("note", "")
    return out


def summary_line(info: Dict[str, Any]) -> str:
    kb = info.get("size", 0) / 1024
    size = f"{kb / 1024:.1f} MB" if kb >= 1024 else f"{kb:.0f} KB"
    return f"«{info.get('name', '')}» · {info.get('label', '')} · {size}"


def describe_for_model(infos: List[Dict[str, Any]], budget: int = 9000) -> List[str]:
    """Un bloque por archivo, repartiendo el espacio: el modelo local tiene poco contexto."""
    readable = [i for i in infos if i.get("text")]
    share = max(1500, budget // max(1, len(readable)))
    blocks = []
    for info in infos:
        text = info.get("text") or ""
        head = summary_line(info)
        where = f"Guardado en workspace/{info['saved']}" if info.get("saved") else ""
        if len(text) > share or info.get("truncated"):
            full = f"; texto completo en workspace/{info['text_file']} (léelo con read_workspace_file y offset)" \
                if info.get("text_file") else ""
            body = text[:share] + f"\n… (recortado{full})"
        else:
            body = text
        notes = " ".join(x for x in (info.get("note", ""), where) if x)
        blocks.append(f"{head}\n{notes}\n{body}".strip())
    return blocks
